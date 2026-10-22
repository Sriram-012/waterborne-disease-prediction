import pptx

template_path = r"C:\Users\LENOVO\.gemini\antigravity\brain\38e0aa45-c2da-4a11-a610-4a454595a8fa\.user_uploaded\media_1790392694666.pptx"
prs = pptx.Presentation(template_path)

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s in slide.shapes:
        if s.has_text_frame:
            print(f"  [Shape ID {s.shape_id}] '{s.name}' (left={s.left}, top={s.top}, w={s.width}, h={s.height}):")
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"      Paragraph: {txt}")
        elif s.has_table:
            print(f"  [Table ID {s.shape_id}] '{s.name}' ({len(s.table.rows)} rows x {len(s.table.columns)} cols):")
            for r_idx, row in enumerate(s.table.rows):
                row_vals = [c.text.strip().replace('\n', ' ') for c in row.cells]
                print(f"      Row {r_idx}: {row_vals}")
